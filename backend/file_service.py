from __future__ import annotations

import io
import os
import re
from pathlib import Path

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_FILE_CHARS = 6000
MAX_PDF_PAGES = 40

TEXT_EXT = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".css", ".js", ".jsx", ".ts",
    ".py", ".java", ".c", ".cpp", ".cs", ".sql", ".log", ".yml", ".yaml",
}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}


class FileError(Exception):
    def __init__(self, code: str, status: int) -> None:
        super().__init__(code)
        self.code = code
        self.status = status


def _finish(text: str) -> tuple[str, bool]:
    text = text.replace("\r\n", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise FileError("empty", 422)
    return text[:MAX_FILE_CHARS], len(text) > MAX_FILE_CHARS


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        parts: list[str] = []
        size = 0
        for page in reader.pages[:MAX_PDF_PAGES]:
            chunk = page.extract_text() or ""
            parts.append(chunk)
            size += len(chunk)
            if size > MAX_FILE_CHARS * 2:
                break
    except Exception as exc:
        raise FileError("failed", 422) from exc
    return "\n".join(parts)


def _docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    except Exception as exc:
        raise FileError("failed", 422) from exc
    return "\n".join(parts)


_WINDOWS_TESSERACT = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
)


def _find_tesseract() -> str | None:
    for candidate in _WINDOWS_TESSERACT:
        if os.path.isfile(candidate):
            return candidate
    return None


def _image(data: bytes) -> str:
    try:
        import pytesseract
        from PIL import Image, ImageOps
    except ImportError as exc:
        raise FileError("ocr_unavailable", 501) from exc

    command = os.environ.get("FLARE_TESSERACT_CMD") or _find_tesseract()
    if command:
        pytesseract.pytesseract.tesseract_cmd = command
    try:
        installed = set(pytesseract.get_languages(config=""))
    except Exception as exc:
        raise FileError("ocr_unavailable", 501) from exc

    wanted = [code for code in ("aze", "rus", "eng") if code in installed]
    lang = "+".join(wanted) if wanted else "eng"
    try:
        image = Image.open(io.BytesIO(data))
        image.load()
        image = ImageOps.autocontrast(ImageOps.exif_transpose(image).convert("L"))
        longest = max(image.size)
        if longest > 3000:
            image.thumbnail((3000, 3000))
        elif longest < 1200:
            scale = 1200 / longest
            image = image.resize((int(image.width * scale), int(image.height * scale)))
        info = pytesseract.image_to_data(image, lang=lang, output_type=pytesseract.Output.DICT)
    except Exception as exc:
        raise FileError("failed", 422) from exc

    lines: dict[tuple, list[str]] = {}
    for i, word in enumerate(info["text"]):
        word = word.strip()
        try:
            conf = float(info["conf"][i])
        except ValueError:
            continue
        if not word or conf < 50:
            continue
        key = (info["block_num"][i], info["par_num"][i], info["line_num"][i])
        lines.setdefault(key, []).append(word)
    text = "\n".join(" ".join(words) for words in lines.values())
    if len(re.findall(r"\w", text)) < 8:
        return ""
    return text


def extract_text(name: str, data: bytes) -> tuple[str, bool]:
    ext = Path(name).suffix.lower()
    if ext in TEXT_EXT:
        text = data.decode("utf-8-sig", errors="replace")
    elif ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    elif ext in IMAGE_EXT:
        text = _image(data)
    else:
        raise FileError("unsupported", 415)
    return _finish(text)